#!/usr/bin/env python3
"""Convert a PPT-safe SVG into editable PowerPoint shapes on one slide.

The generated slide XML is transplanted into the original package so every
non-target PPTX part keeps its original uncompressed bytes.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - environment diagnostic
    raise SystemExit(
        "python-pptx is required. Use the bundled Codex workspace Python or "
        "install python-pptx>=1.0."
    ) from exc


SVG_NS = "http://www.w3.org/2000/svg"
NS = {"svg": SVG_NS}
NAMED_COLORS = {
    "white": "FFFFFF",
    "black": "000000",
    "red": "FF0000",
    "green": "008000",
    "blue": "0000FF",
    "gray": "808080",
    "grey": "808080",
    "none": None,
    "transparent": None,
}
FILL_TAGS = {"solidFill", "gradFill", "noFill", "pattFill", "blipFill", "grpFill"}
SUPPORTED_RENDER_TAGS = {"svg", "g", "rect", "circle", "ellipse", "line", "path", "polygon", "polyline", "text", "tspan"}
IGNORED_DEF_TAGS = {
    "defs",
    "style",
    "title",
    "desc",
    "linearGradient",
    "stop",
    "marker",
    "filter",
    "feDropShadow",
}
UNSUPPORTED_TAGS = {"image", "use", "foreignObject", "textPath", "clipPath", "mask", "pattern", "symbol"}
PATH_TOKEN_RE = re.compile(r"[MmLlHhVvQqCcZz]|[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?")
LENGTH_RE = re.compile(r"[-+]?(?:\d*\.\d+|\d+\.?)")


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def parse_number(value: str | None, default: float = 0.0) -> float:
    if not value:
        return default
    match = LENGTH_RE.search(value)
    return float(match.group(0)) if match else default


def normalize_color(value: str | None) -> str | None:
    if value is None:
        return None
    value = value.strip()
    if value.lower() in NAMED_COLORS:
        return NAMED_COLORS[value.lower()]
    value = value.removeprefix("#")
    if re.fullmatch(r"[0-9a-fA-F]{3}", value):
        value = "".join(char * 2 for char in value)
    return value.upper() if re.fullmatch(r"[0-9a-fA-F]{6}", value) else None


def parse_css(root: ET.Element) -> dict[str, dict[str, str]]:
    rules: dict[str, dict[str, str]] = {}
    for style in root.findall(".//svg:style", NS):
        css = "".join(style.itertext())
        css = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
        for selectors, body in re.findall(r"([^{}]+)\{([^{}]+)\}", css):
            props: dict[str, str] = {}
            for declaration in body.split(";"):
                if ":" not in declaration:
                    continue
                key, value = declaration.split(":", 1)
                props[key.strip()] = value.strip()
            for selector in selectors.split(","):
                rules[selector.strip()] = props.copy()
    return rules


GEOMETRY_ATTRS = {
    "class",
    "id",
    "d",
    "points",
    "x",
    "y",
    "x1",
    "y1",
    "x2",
    "y2",
    "width",
    "height",
    "rx",
    "ry",
    "cx",
    "cy",
    "r",
    "dx",
    "dy",
}


def element_style(element: ET.Element, rules: dict[str, dict[str, str]]) -> dict[str, str]:
    result: dict[str, str] = {}
    result.update(rules.get(local_name(element.tag), {}))
    for class_name in element.get("class", "").split():
        result.update(rules.get(f".{class_name}", {}))
    result.update({key: value for key, value in element.attrib.items() if key not in GEOMETRY_ATTRS})
    for declaration in element.get("style", "").split(";"):
        if ":" in declaration:
            key, value = declaration.split(":", 1)
            result[key.strip()] = value.strip()
    return result


@dataclass
class PreflightResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "errors": self.errors,
            "warnings": self.warnings,
            "elements": dict(sorted(self.counts.items())),
        }


def preflight_svg(root: ET.Element) -> PreflightResult:
    result = PreflightResult()
    if local_name(root.tag) != "svg":
        result.errors.append("Root element is not <svg>.")
    width = parse_number(root.get("width"))
    height = parse_number(root.get("height"))
    view_box = root.get("viewBox")
    if not width or not height:
        if view_box:
            values = [parse_number(item) for item in re.split(r"[ ,]+", view_box.strip())]
            if len(values) != 4 or values[2] <= 0 or values[3] <= 0:
                result.errors.append("SVG needs positive width/height or a valid viewBox.")
        else:
            result.errors.append("SVG needs positive width/height or a valid viewBox.")

    seen_filter_warning = False
    for element in root.iter():
        tag = local_name(element.tag)
        result.counts[tag] = result.counts.get(tag, 0) + 1
        if tag in UNSUPPORTED_TAGS:
            result.errors.append(f"Unsupported SVG element: <{tag}>.")
        elif tag not in SUPPORTED_RENDER_TAGS and tag not in IGNORED_DEF_TAGS:
            result.errors.append(f"Unknown SVG element: <{tag}>.")
        if element.get("transform"):
            result.errors.append(f"Transforms are not supported ({tag}: {element.get('transform')}).")
        if tag == "path":
            commands = set(re.findall(r"[A-Za-z]", element.get("d", "")))
            unsupported = sorted(commands - set("MmLlHhVvQqCcZz"))
            if unsupported:
                result.errors.append(f"Unsupported path commands: {', '.join(unsupported)}.")
        if (element.get("filter") or "filter:" in element.get("style", "")) and not seen_filter_warning:
            result.warnings.append("SVG filters are ignored; verify shadows visually in the rendered slide.")
            seen_filter_warning = True
        for attr in ("clip-path", "mask"):
            if element.get(attr):
                result.errors.append(f"Unsupported attribute {attr} on <{tag}>.")
        for attr in ("opacity", "fill-opacity", "stroke-opacity"):
            if element.get(attr) not in (None, "1", "1.0", "100%"):
                result.warnings.append(f"{attr} on <{tag}> may be approximated.")

    result.errors = list(dict.fromkeys(result.errors))
    result.warnings = list(dict.fromkeys(result.warnings))
    return result


def parse_gradients(root: ET.Element) -> dict[str, dict[str, Any]]:
    gradients: dict[str, dict[str, Any]] = {}
    for gradient in root.findall(".//svg:linearGradient", NS):
        stops: list[tuple[float, str, float]] = []
        for stop in gradient.findall("svg:stop", NS):
            style = {}
            for declaration in stop.get("style", "").split(";"):
                if ":" in declaration:
                    key, value = declaration.split(":", 1)
                    style[key.strip()] = value.strip()
            raw_offset = stop.get("offset", "0")
            offset = parse_number(raw_offset) / (100 if "%" in raw_offset else 1)
            rgb = normalize_color(stop.get("stop-color") or style.get("stop-color")) or "FFFFFF"
            opacity = parse_number(stop.get("stop-opacity") or style.get("stop-opacity"), 1)
            stops.append((max(0.0, min(1.0, offset)), rgb, max(0.0, min(1.0, opacity))))
        if not stops:
            continue
        horizontal_delta = abs(parse_number(gradient.get("x2"), 1) - parse_number(gradient.get("x1")))
        vertical_delta = abs(parse_number(gradient.get("y2")) - parse_number(gradient.get("y1")))
        gradients[gradient.get("id", "")] = {
            "stops": stops,
            "vertical": vertical_delta > horizontal_delta,
        }
    return gradients


def set_gradient_fill(shape, gradient: dict[str, Any]) -> None:
    sp_pr = shape._element.spPr
    for child in list(sp_pr):
        if local_name(child.tag) in FILL_TAGS:
            sp_pr.remove(child)
    grad_fill = OxmlElement("a:gradFill")
    grad_fill.set("rotWithShape", "1")
    gs_list = OxmlElement("a:gsLst")
    for position, rgb, opacity in gradient["stops"]:
        gs = OxmlElement("a:gs")
        gs.set("pos", str(round(position * 100000)))
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", rgb)
        if opacity < 0.999:
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(round(opacity * 100000)))
            srgb.append(alpha)
        gs.append(srgb)
        gs_list.append(gs)
    grad_fill.append(gs_list)
    linear = OxmlElement("a:lin")
    linear.set("ang", "5400000" if gradient["vertical"] else "0")
    linear.set("scaled", "1")
    grad_fill.append(linear)
    insert_at = len(sp_pr)
    for index, child in enumerate(sp_pr):
        if local_name(child.tag) in {"ln", "effectLst", "effectDag", "scene3d", "sp3d"}:
            insert_at = index
            break
    sp_pr.insert(insert_at, grad_fill)


def set_solid_alpha(shape, opacity: float) -> None:
    if opacity >= 0.999:
        return
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None or not len(solid):
        return
    color_node = solid[0]
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(round(max(0.0, min(1.0, opacity)) * 100000)))
    color_node.append(alpha)


def apply_fill(shape, value: str | None, style: dict[str, str], gradients: dict[str, dict[str, Any]]) -> None:
    if value and value.startswith("url(#"):
        gradient = gradients.get(value[5:-1])
        if gradient:
            set_gradient_fill(shape, gradient)
            return
    rgb = normalize_color(value)
    if rgb is None:
        shape.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(rgb)
    opacity = parse_number(style.get("fill-opacity") or style.get("opacity"), 1)
    set_solid_alpha(shape, opacity)


def apply_line(shape, style: dict[str, str], scale_in: float) -> None:
    rgb = normalize_color(style.get("stroke"))
    if rgb is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGBColor.from_string(rgb)
    shape.line.width = Pt(max(0.35, parse_number(style.get("stroke-width"), 1) * scale_in * 72))
    if style.get("stroke-dasharray"):
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def add_arrowheads(shape, style: dict[str, str]) -> None:
    line = shape._element.spPr.find(qn("a:ln"))
    if line is None:
        line = OxmlElement("a:ln")
        shape._element.spPr.append(line)
    if style.get("marker-end"):
        head = OxmlElement("a:headEnd")
        head.set("type", "triangle")
        head.set("w", "sm")
        head.set("len", "sm")
        line.append(head)
    if style.get("marker-start"):
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        tail.set("w", "sm")
        tail.set("len", "sm")
        line.append(tail)


def parse_points(raw: str) -> list[tuple[float, float]]:
    values = [float(value) for value in re.findall(r"[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?", raw)]
    if len(values) % 2:
        raise ValueError("points attribute contains an odd number of coordinates")
    return list(zip(values[::2], values[1::2]))


def parse_path(data: str) -> tuple[list[tuple[float, float]], bool]:
    tokens = PATH_TOKEN_RE.findall(data)
    points: list[tuple[float, float]] = []
    index = 0
    command = ""
    current = (0.0, 0.0)
    start = current
    closed = False

    def take() -> float:
        nonlocal index
        value = float(tokens[index])
        index += 1
        return value

    def absolute(x: float, y: float, relative: bool) -> tuple[float, float]:
        return (current[0] + x, current[1] + y) if relative else (x, y)

    while index < len(tokens):
        if re.fullmatch(r"[MmLlHhVvQqCcZz]", tokens[index]):
            command = tokens[index]
            index += 1
            if command in "Zz":
                closed = True
                current = start
                continue
        relative = command.islower()
        upper = command.upper()
        if upper == "M":
            current = absolute(take(), take(), relative)
            start = current
            points.append(current)
            command = "l" if relative else "L"
        elif upper == "L":
            current = absolute(take(), take(), relative)
            points.append(current)
        elif upper == "H":
            value = take()
            current = (current[0] + value, current[1]) if relative else (value, current[1])
            points.append(current)
        elif upper == "V":
            value = take()
            current = (current[0], current[1] + value) if relative else (current[0], value)
            points.append(current)
        elif upper == "Q":
            p0 = current
            control = absolute(take(), take(), relative)
            end = absolute(take(), take(), relative)
            for step in range(1, 9):
                t = step / 8
                u = 1 - t
                points.append((
                    u * u * p0[0] + 2 * u * t * control[0] + t * t * end[0],
                    u * u * p0[1] + 2 * u * t * control[1] + t * t * end[1],
                ))
            current = end
        elif upper == "C":
            p0 = current
            c1 = absolute(take(), take(), relative)
            c2 = absolute(take(), take(), relative)
            end = absolute(take(), take(), relative)
            for step in range(1, 11):
                t = step / 10
                u = 1 - t
                points.append((
                    u**3 * p0[0] + 3 * u * u * t * c1[0] + 3 * u * t * t * c2[0] + t**3 * end[0],
                    u**3 * p0[1] + 3 * u * u * t * c1[1] + 3 * u * t * t * c2[1] + t**3 * end[1],
                ))
            current = end
        else:
            break
    return points, closed


def estimated_text_width(text: str, font_size: float) -> float:
    return max(
        font_size * 1.8,
        sum(font_size * (1.0 if ord(char) > 255 else 0.58) for char in text) + font_size * 0.6,
    )


@dataclass
class RenderOptions:
    x: float | None
    y: float
    max_width: float
    max_height: float
    font: str
    min_font_size: float
    prefix: str


class SvgRenderer:
    def __init__(self, slide, root: ET.Element, options: RenderOptions, slide_width_in: float):
        self.slide = slide
        self.root = root
        self.rules = parse_css(root)
        self.gradients = parse_gradients(root)
        view_box = root.get("viewBox")
        if view_box:
            values = [parse_number(item) for item in re.split(r"[ ,]+", view_box.strip())]
            self.view_x, self.view_y, self.svg_width, self.svg_height = values
        else:
            self.view_x = self.view_y = 0.0
            self.svg_width = parse_number(root.get("width"))
            self.svg_height = parse_number(root.get("height"))
        self.scale_in = min(options.max_width / self.svg_width, options.max_height / self.svg_height)
        self.x0_in = options.x if options.x is not None else (slide_width_in - self.svg_width * self.scale_in) / 2
        self.y0_in = options.y
        self.options = options
        self.count = 0

    def x(self, value: float):
        return Inches(self.x0_in + (value - self.view_x) * self.scale_in)

    def y(self, value: float):
        return Inches(self.y0_in + (value - self.view_y) * self.scale_in)

    def length(self, value: float):
        return Inches(value * self.scale_in)

    def name_shape(self, shape, kind: str, detail: str = "") -> None:
        self.count += 1
        clean = re.sub(r"\s+", " ", detail).strip()[:28]
        shape.name = f"{self.options.prefix}-{self.count:04d}-{kind}" + (f"-{clean}" if clean else "")

    def add_text(self, element: ET.Element, style: dict[str, str]) -> None:
        tspans = ["".join(child.itertext()).strip() for child in list(element) if local_name(child.tag) == "tspan"]
        lines = [line for line in tspans if line] or ["".join(element.itertext()).strip()]
        if not lines[0]:
            return
        x = parse_number(element.get("x"))
        y = parse_number(element.get("y"))
        font_px = parse_number(style.get("font-size"), 16)
        font_pt = max(self.options.min_font_size, font_px * self.scale_in * 72)
        width_units = max(estimated_text_width(line, font_px) for line in lines)
        anchor = style.get("text-anchor", "start")
        if anchor == "middle":
            left, alignment = x - width_units / 2, PP_ALIGN.CENTER
        elif anchor == "end":
            left, alignment = x - width_units, PP_ALIGN.RIGHT
        else:
            left, alignment = x, PP_ALIGN.LEFT
        top = y - font_px * 0.96
        height = max(font_px * 1.32 * len(lines), 18)
        box = self.slide.shapes.add_textbox(self.x(left), self.y(top), self.length(width_units), self.length(height))
        self.name_shape(box, "text", " / ".join(lines))
        frame = box.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        paragraph = frame.paragraphs[0]
        paragraph.alignment = alignment
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.text = "\v".join(lines)
        paragraph.font.name = self.options.font
        paragraph.font.size = Pt(font_pt)
        paragraph.font.bold = parse_number(style.get("font-weight"), 400) >= 650
        paragraph.font.italic = style.get("font-style") == "italic"
        paragraph.font.color.rgb = RGBColor.from_string(normalize_color(style.get("fill")) or "173B56")

    def add_rect(self, element: ET.Element, style: dict[str, str]) -> None:
        x, y = parse_number(element.get("x")), parse_number(element.get("y"))
        width, height = parse_number(element.get("width")), parse_number(element.get("height"))
        if (
            math.isclose(x, self.view_x)
            and math.isclose(y, self.view_y)
            and math.isclose(width, self.svg_width)
            and math.isclose(height, self.svg_height)
            and normalize_color(style.get("fill")) == "FFFFFF"
        ):
            return
        radius = max(parse_number(element.get("rx")), parse_number(element.get("ry")))
        shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = self.slide.shapes.add_shape(shape_type, self.x(x), self.y(y), self.length(width), self.length(height))
        if radius and min(width, height) > 0:
            shape.adjustments[0] = min(0.5, radius / min(width, height))
        self.name_shape(shape, "rect", element.get("id", ""))
        apply_fill(shape, style.get("fill"), style, self.gradients)
        apply_line(shape, style, self.scale_in)

    def add_ellipse(self, element: ET.Element, style: dict[str, str], circle: bool) -> None:
        cx, cy = parse_number(element.get("cx")), parse_number(element.get("cy"))
        rx = parse_number(element.get("r")) if circle else parse_number(element.get("rx"))
        ry = parse_number(element.get("r")) if circle else parse_number(element.get("ry"))
        shape = self.slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL,
            self.x(cx - rx),
            self.y(cy - ry),
            self.length(2 * rx),
            self.length(2 * ry),
        )
        self.name_shape(shape, "circle" if circle else "ellipse", element.get("id", ""))
        apply_fill(shape, style.get("fill"), style, self.gradients)
        apply_line(shape, style, self.scale_in)

    def add_line(self, element: ET.Element, style: dict[str, str]) -> None:
        shape = self.slide.shapes.add_connector(
            MSO_CONNECTOR_TYPE.STRAIGHT,
            self.x(parse_number(element.get("x1"))),
            self.y(parse_number(element.get("y1"))),
            self.x(parse_number(element.get("x2"))),
            self.y(parse_number(element.get("y2"))),
        )
        self.name_shape(shape, "line", element.get("id", ""))
        apply_line(shape, style, self.scale_in)
        add_arrowheads(shape, style)

    def add_freeform(self, points: list[tuple[float, float]], closed: bool, style: dict[str, str], kind: str, detail: str = "") -> None:
        if len(points) < 2:
            return
        builder = self.slide.shapes.build_freeform(
            points[0][0] - self.view_x,
            points[0][1] - self.view_y,
            scale=Inches(self.scale_in),
        )
        shifted = [(x - self.view_x, y - self.view_y) for x, y in points[1:]]
        builder.add_line_segments(shifted, close=closed)
        shape = builder.convert_to_shape(origin_x=Inches(self.x0_in), origin_y=Inches(self.y0_in))
        self.name_shape(shape, kind, detail)
        apply_fill(shape, style.get("fill"), style, self.gradients)
        apply_line(shape, style, self.scale_in)
        add_arrowheads(shape, style)

    def walk(self, parent: ET.Element) -> None:
        for element in list(parent):
            tag = local_name(element.tag)
            if tag in IGNORED_DEF_TAGS:
                continue
            style = element_style(element, self.rules)
            if tag == "g":
                self.walk(element)
            elif tag == "rect":
                self.add_rect(element, style)
            elif tag == "circle":
                self.add_ellipse(element, style, True)
            elif tag == "ellipse":
                self.add_ellipse(element, style, False)
            elif tag == "line":
                self.add_line(element, style)
            elif tag == "path":
                points, closed = parse_path(element.get("d", ""))
                self.add_freeform(points, closed, style, "path", element.get("id", ""))
            elif tag in {"polygon", "polyline"}:
                self.add_freeform(parse_points(element.get("points", "")), tag == "polygon", style, tag, element.get("id", ""))
            elif tag == "text":
                self.add_text(element, style)

    def render(self) -> int:
        self.walk(self.root)
        return self.count


def slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))


def clear_slide(slide, mode: str, prefix: str, keep_placeholders: bool) -> int:
    tree = slide.shapes._spTree
    removed = 0
    for shape in list(slide.shapes):
        if mode == "append":
            continue
        if mode == "replace-tagged" and not shape.name.startswith(prefix):
            continue
        if mode == "replace-slide" and keep_placeholders and shape.is_placeholder:
            continue
        tree.remove(shape._element)
        removed += 1
    return removed


def write_minimal_patch(original: Path, generated: Path, output: Path, target_part: str) -> dict[str, Any]:
    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(original, "r") as source, zipfile.ZipFile(generated, "r") as changed:
        if target_part not in source.namelist() or target_part not in changed.namelist():
            raise RuntimeError(f"Target slide part not found: {target_part}")
        with tempfile.NamedTemporaryFile(prefix="svg2pptx-", suffix=".pptx", dir=output.parent, delete=False) as handle:
            temp_output = Path(handle.name)
        try:
            with zipfile.ZipFile(temp_output, "w") as destination:
                for info in source.infolist():
                    data = changed.read(info.filename) if info.filename == target_part else source.read(info.filename)
                    destination.writestr(info, data)
            os.replace(temp_output, output)
        finally:
            temp_output.unlink(missing_ok=True)

    with zipfile.ZipFile(original, "r") as before, zipfile.ZipFile(output, "r") as after:
        changed_parts = [name for name in before.namelist() if before.read(name) != after.read(name)]
        extra_parts = sorted(set(after.namelist()) - set(before.namelist()))
        missing_parts = sorted(set(before.namelist()) - set(after.namelist()))
    if changed_parts != [target_part] or extra_parts or missing_parts:
        raise RuntimeError(
            "Minimal-patch invariant failed: "
            f"changed={changed_parts}, extra={extra_parts}, missing={missing_parts}"
        )
    return {
        "changed_parts": changed_parts,
        "unchanged_parts": len(before.namelist()) - 1,
    }


def backup_path_for(path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return path.with_name(f"{path.stem}.before-svg2pptx-{timestamp}{path.suffix}")


def positive_float(value: str) -> float:
    parsed = float(value)
    if parsed <= 0:
        raise argparse.ArgumentTypeError("value must be positive")
    return parsed


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--svg", required=True, type=Path, help="PPT-safe SVG input")
    parser.add_argument("--preflight-only", action="store_true", help="Only inspect SVG compatibility")
    parser.add_argument("--fail-on-warning", action="store_true", help="Treat preflight warnings as errors")
    parser.add_argument("--pptx", type=Path, help="Existing PowerPoint presentation")
    parser.add_argument("--slide", type=int, help="1-based target slide index")
    destination = parser.add_mutually_exclusive_group()
    destination.add_argument("--output", type=Path, help="Write to a new PPTX")
    destination.add_argument("--in-place", action="store_true", help="Update --pptx and create a timestamped backup")
    parser.add_argument("--mode", choices=("append", "replace-tagged", "replace-slide"), default="append")
    parser.add_argument("--expected-text", help="Require this text on the target slide before editing")
    parser.add_argument("--force", action="store_true", help="Allow replace-slide without expected-text")
    parser.add_argument("--drop-placeholders", action="store_true", help="Remove placeholders in replace-slide mode")
    parser.add_argument("--prefix", default="svg2pptx", help="Name prefix for generated PowerPoint shapes")
    parser.add_argument("--x", type=float, help="Left position in inches; default centers horizontally")
    parser.add_argument("--y", type=float, default=0.5, help="Top position in inches")
    parser.add_argument("--max-width", type=positive_float, help="Maximum width in inches")
    parser.add_argument("--max-height", type=positive_float, help="Maximum height in inches")
    parser.add_argument("--font", default="Microsoft YaHei", help="PowerPoint font for converted text")
    parser.add_argument("--min-font-size", type=positive_float, default=6.8, help="Minimum text size in points")
    parser.add_argument("--overwrite", action="store_true", help="Overwrite an existing --output file")
    parser.add_argument("--report-json", type=Path, help="Also write the JSON report to this path")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    if not args.svg.is_file():
        raise SystemExit(f"SVG not found: {args.svg}")
    root = ET.parse(args.svg).getroot()
    preflight = preflight_svg(root)
    report: dict[str, Any] = {"svg": str(args.svg.resolve()), "preflight": preflight.as_dict()}
    if args.preflight_only:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if preflight.ok and not (args.fail_on_warning and preflight.warnings) else 2
    if not preflight.ok or (args.fail_on_warning and preflight.warnings):
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 2
    if not args.pptx or not args.pptx.is_file():
        raise SystemExit("--pptx must reference an existing presentation")
    if not args.slide or args.slide < 1:
        raise SystemExit("--slide must be a positive 1-based index")
    if not args.in_place and not args.output:
        raise SystemExit("Choose --output or --in-place")
    if args.mode == "replace-slide" and not args.expected_text and not args.force:
        raise SystemExit("replace-slide requires --expected-text or explicit --force")

    input_pptx = args.pptx.resolve()
    output_pptx = input_pptx if args.in_place else args.output.resolve()
    if not args.in_place and output_pptx.exists() and not args.overwrite:
        raise SystemExit(f"Output exists; pass --overwrite: {output_pptx}")

    backup = None
    if args.in_place:
        backup = backup_path_for(input_pptx)
        shutil.copy2(input_pptx, backup)

    try:
        with tempfile.TemporaryDirectory(prefix="svg2pptx-") as temp_dir:
            generated = Path(temp_dir) / "generated.pptx"
            shutil.copy2(input_pptx, generated)
            presentation = Presentation(generated)
            if args.slide > len(presentation.slides):
                raise RuntimeError(f"Slide {args.slide} does not exist; deck has {len(presentation.slides)} slides")
            slide = presentation.slides[args.slide - 1]
            existing_text = slide_text(slide)
            if args.expected_text and args.expected_text not in existing_text:
                raise RuntimeError(f"Target slide does not contain expected text: {args.expected_text!r}")
            removed = clear_slide(slide, args.mode, args.prefix, not args.drop_placeholders)
            if args.mode == "replace-tagged" and removed == 0:
                raise RuntimeError(f"No shapes matched prefix {args.prefix!r}")
            slide_width = presentation.slide_width / 914400
            slide_height = presentation.slide_height / 914400
            max_width = args.max_width or max(0.1, slide_width - 1.0)
            max_height = args.max_height or max(0.1, slide_height - args.y - 0.5)
            if args.x is not None and (args.x < 0 or args.x + max_width > slide_width + 1e-6):
                raise RuntimeError("Requested x/max-width exceeds slide bounds")
            if args.y < 0 or args.y + max_height > slide_height + 1e-6:
                raise RuntimeError("Requested y/max-height exceeds slide bounds")
            options = RenderOptions(
                x=args.x,
                y=args.y,
                max_width=max_width,
                max_height=max_height,
                font=args.font,
                min_font_size=args.min_font_size,
                prefix=args.prefix,
            )
            renderer = SvgRenderer(slide, root, options, slide_width)
            shape_count = renderer.render()
            target_part = str(slide.part.partname).lstrip("/")
            presentation.save(generated)
            patch_source = backup if args.in_place else input_pptx
            patch_report = write_minimal_patch(patch_source, generated, output_pptx, target_part)
    except Exception:
        if args.in_place and backup and backup.exists():
            shutil.copy2(backup, input_pptx)
        raise

    report.update(
        {
            "pptx": str(input_pptx),
            "output": str(output_pptx),
            "backup": str(backup) if backup else None,
            "slide": args.slide,
            "mode": args.mode,
            "removed_shapes": removed,
            "created_shapes": shape_count,
            "placement": {
                "x": args.x if args.x is not None else "center",
                "y": args.y,
                "max_width": max_width,
                "max_height": max_height,
            },
            "package": patch_report,
        }
    )
    output = json.dumps(report, ensure_ascii=False, indent=2)
    print(output)
    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(output + "\n", encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
